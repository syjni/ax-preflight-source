from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ax_scanner.models import FileRecord, ScanReport, TableProfile

from .errors import AgentToolError


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


class TableStore:
    """Server-side row loader keyed entirely by the stable scanner report."""

    def __init__(self, report: ScanReport, source_root: Path):
        self.source_root = source_root.resolve()
        self.files = {record.file_id: record for record in report.files}
        self.tables = {table.table_id: table for table in report.tables}
        self.cache: dict[str, list[dict[str, Any]]] = {}

    def profile(self, table_id: str) -> TableProfile:
        try:
            return self.tables[table_id]
        except KeyError as exc:
            raise AgentToolError("TABLE_NOT_FOUND", f"unknown table_id: {table_id}") from exc

    def rows(self, table_id: str) -> list[dict[str, Any]]:
        if table_id not in self.cache:
            profile = self.profile(table_id)
            record = self.files.get(profile.file_id)
            if record is None:
                raise AgentToolError("SCANNER_CONTRACT_MISMATCH", f"missing file for {table_id}")
            path = (self.source_root / Path(record.relative_path)).resolve()
            if not path.is_relative_to(self.source_root):
                raise AgentToolError("UNSAFE_SOURCE_PATH", f"table path escapes source root: {record.relative_path}")
            if not path.is_file():
                raise AgentToolError("SOURCE_FILE_MISSING", f"source file not found: {record.relative_path}")
            if record.extension.lower() == ".csv":
                matrix = self._read_csv(path, record)
            elif record.extension.lower() == ".xlsx":
                matrix = self._read_xlsx(path, profile.sheet_name)
            else:
                raise AgentToolError("UNSUPPORTED_TABLE_SOURCE", f"unsupported table source: {record.extension}")
            self.cache[table_id] = self._rows_from_matrix(matrix, profile)
        return self.cache[table_id]

    @staticmethod
    def _read_csv(path: Path, record: FileRecord) -> list[list[Any]]:
        encoding = str(record.parser_metadata.get("encoding") or "utf-8-sig")
        delimiter = str(record.parser_metadata.get("delimiter") or ",")
        try:
            text = path.read_text(encoding=encoding)
        except (LookupError, UnicodeDecodeError) as exc:
            raise AgentToolError("TABLE_READ_ERROR", f"cannot decode {record.relative_path}") from exc
        return [list(row) for row in csv.reader(text.splitlines(), delimiter=delimiter)]

    @staticmethod
    def _read_xlsx(path: Path, sheet_name: str) -> list[list[Any]]:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise AgentToolError("SCANNER_CONTRACT_MISMATCH", f"sheet missing from source: {sheet_name}")
            return [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
        finally:
            workbook.close()

    @staticmethod
    def _rows_from_matrix(matrix: list[list[Any]], profile: TableProfile) -> list[dict[str, Any]]:
        first_nonempty = next(
            (index for index, row in enumerate(matrix) if any(not _is_empty(value) for value in row)),
            None,
        )
        data_rows = [] if first_nonempty is None else [
            row for row in matrix[first_nonempty + 1 :] if any(not _is_empty(value) for value in row)
        ]
        width = len(profile.column_names)
        materialized = [
            dict(zip(profile.column_names, row[:width] + [None] * max(0, width - len(row))))
            for row in data_rows
        ]
        if len(materialized) != profile.row_count:
            raise AgentToolError(
                "SCANNER_CONTRACT_MISMATCH",
                f"{profile.table_id} row count changed: report={profile.row_count}, source={len(materialized)}",
            )
        return materialized
