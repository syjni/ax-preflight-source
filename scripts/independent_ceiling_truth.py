"""Independently derive Ceiling answers from generated company evidence.

This module never opens benchmark_tasks.json and therefore cannot reuse a
declared benchmark answer while computing source truth.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


def _csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _xlsx_rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = [str(value) for value in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]


def _docx_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _number_from(text: str, pattern: str) -> int:
    match = re.search(pattern, text)
    if not match:
        raise ValueError(f"source text did not match {pattern!r}")
    return int(match.group(1))


def independently_compute_truth(source_root: Path) -> dict[str, Any]:
    """Calculate every answer from runtime-facing source files alone."""
    policy_root = source_root / "02_정책"
    sales_root = source_root / "03_영업"
    order_root = source_root / "04_주문"
    logistics_root = source_root / "05_물류"
    support_root = source_root / "06_고객지원"

    return_policy = _docx_text(policy_root / "반품_교환_정책_2026.docx")
    delay_policy = _docx_text(policy_root / "배송_지연_보상_2026.docx")
    discount_policy = _docx_text(policy_root / "할인_승인_규정.docx")
    warehouse_manual = _docx_text(logistics_root / "창고_운영_매뉴얼.docx")
    support_manual = _docx_text(support_root / "고객응대_매뉴얼.docx")

    customers = _xlsx_rows(sales_root / "거래처_마스터.xlsx", "Customers")
    products = _csv_rows(sales_root / "품목_마스터.csv")
    orders = _xlsx_rows(order_root / "주문_원장_2026.xlsx", "Orders")
    returns = _csv_rows(order_root / "반품_원장_2026.csv")
    inventory = _xlsx_rows(logistics_root / "재고_현황_2026-09.xlsx", "Inventory")
    lead_times = _csv_rows(logistics_root / "배송권역_리드타임.csv")

    customer_ids = {str(row["customer_name"]): str(row["customer_id"]) for row in customers}
    product_ids = {str(row["product_name"]): str(row["product_id"]) for row in products}
    integer = lambda value: int(value)  # noqa: E731 - keeps source calculations legible.
    by_customer = lambda name: customer_ids[name]  # noqa: E731

    return_requirements = [
        item for item in ("미개봉", "미사용", "재판매 가능 상태", "원본 영수증")
        if item in return_policy
    ]
    if len(return_requirements) != 4:
        raise ValueError("return policy requirements are incomplete in source evidence")

    top_stock = max(inventory, key=lambda row: (integer(row["available_qty"]), str(row["product_id"])))
    daon_orders = [row for row in orders if str(row["customer_id"]) == by_customer("다온유통")]
    daon_top = max(daon_orders, key=lambda row: (integer(row["amount"]), str(row["order_id"])))

    answer_map: dict[str, Any] = {
        "K01_RETURN_WINDOW": _number_from(return_policy, r"(\d+)일"),
        "K02_RETURN_REQUIREMENTS": return_requirements,
        "K03_DELAY_COMPENSATION_THRESHOLD": _number_from(delay_policy, r"(\d+)영업일"),
        "K04_DISCOUNT_APPROVER": "영업운영팀장" if "영업운영팀장" in discount_policy else None,
        "K05_STOCK_REFERENCE_FIELD": "available_qty" if "available_qty" in warehouse_manual else None,
        "K06_FIRST_RESPONSE_TARGET": _number_from(support_manual, r"(\d+)시간"),
        "O01_CUSTOMER_COUNT": len(customers),
        "O02_SEPTEMBER_ORDER_TOTAL": sum(integer(row["amount"]) for row in orders if str(row["order_date"]).startswith("2026-09")),
        "O03_GREEN_TABLE_LAST_ORDER": max(str(row["order_date"]) for row in orders if str(row["customer_id"]) == by_customer("그린테이블")),
        "O04_PRODUCT_MASTER_COUNT": len(products),
        "O05_HIGHEST_AVAILABLE_STOCK": str(top_stock["product_id"]),
        "O06_RETURN_RATE": sum(integer(row["returned_qty"]) for row in returns) / sum(integer(row["shipped_qty"]) for row in returns),
        "O07_SEOUL_LEAD_TIME": integer(next(row["standard_lead_days"] for row in lead_times if row["zone_name"] == "서울")),
        "C01_HANBIT_AUGUST_TOTAL": sum(integer(row["amount"]) for row in orders if str(row["customer_id"]) == by_customer("한빛상사") and str(row["order_date"]).startswith("2026-08")),
        "C02_DONGHAE_SEPTEMBER_TOTAL": sum(integer(row["amount"]) for row in orders if str(row["customer_id"]) == by_customer("동해상사") and str(row["order_date"]).startswith("2026-09")),
        "C03_TEA_STOCK": sum(integer(row["available_qty"]) for row in inventory if str(row["product_id"]) == product_ids["프리미엄 홍차 세트"]),
        "C04_YUNSEONG_LAST_ORDER": max(str(row["order_date"]) for row in orders if str(row["customer_id"]) == by_customer("윤성마트")),
        "C05_SAEBOM_JULY_TOTAL": sum(integer(row["amount"]) for row in orders if str(row["customer_id"]) == by_customer("새봄마트") and str(row["order_date"]).startswith("2026-07")),
        "C06_ARAM_AUGUST_RETURNS": sum(integer(row["returned_qty"]) for row in returns if str(row["customer_id"]) == by_customer("아람리빙") and str(row["return_date"]).startswith("2026-08")),
        "C07_DAON_TOP_ORDER_PRODUCT": str(daon_top["product_id"]),
        "K07_VENDOR_FORECAST_ABSTAIN": None,
    }
    return {
        "answers": answer_map,
        "structured_row_counts": {
            "customers": len(customers), "products": len(products), "orders": len(orders),
            "returns": len(returns), "inventory": len(inventory),
        },
        "method": "independent_runtime_source_readers",
    }
