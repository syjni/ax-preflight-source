from __future__ import annotations

import argparse
import os
import re
import shutil
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


FIXED_TIMESTAMP = datetime(2026, 9, 19, tzinfo=timezone.utc).timestamp()
FIXED_DOCUMENT_TIME = datetime(2026, 9, 19, tzinfo=timezone.utc)
ZIP_TIME = (2026, 9, 19, 0, 0, 0)


def _canonicalize_zip(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source:
        members = []
        for info in source.infolist():
            if info.is_dir():
                continue
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2026-09-19T00:00:00Z\g<2>",
                    data,
                )
            members.append((info.filename, data))
    with tempfile.NamedTemporaryFile(delete=False, dir=path.parent, suffix=".zip") as handle:
        temp_path = Path(handle.name)
    try:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
            for name, data in sorted(members):
                info = zipfile.ZipInfo(name, ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                target.writestr(info, data)
        temp_path.replace(path)
    finally:
        temp_path.unlink(missing_ok=True)


def _write_text_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter, invariant=1, pageCompression=0)
    pdf.drawString(72, 720, "Shipping delay compensation starts after 3 days.")
    pdf.drawString(72, 700, "Compensation is issued as store credit.")
    pdf.save()


def _write_blank_pdf(path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as stream:
        writer.write(stream)


def _write_docx(path: Path) -> None:
    document = Document()
    document.core_properties.created = FIXED_DOCUMENT_TIME
    document.core_properties.modified = FIXED_DOCUMENT_TIME
    document.add_paragraph("Warehouse operating guide")
    document.add_paragraph("Use the approved customer identifier before aggregating orders.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Code"
    table.cell(0, 1).text = "Meaning"
    table.cell(1, 0).text = "C013"
    table.cell(1, 1).text = "Hanbit Trading"
    document.save(path)
    _canonicalize_zip(path)


def _write_xlsx(path: Path) -> None:
    workbook = Workbook()
    workbook.properties.created = FIXED_DOCUMENT_TIME.replace(tzinfo=None)
    workbook.properties.modified = FIXED_DOCUMENT_TIME.replace(tzinfo=None)
    orders = workbook.active
    orders.title = "Orders"
    orders.append(["order_id", "customer_id", "order_date", "amount", "memo"])
    orders.append(["O100", "C013", "2026-08-03", 125000, None])
    orders.append(["O101", "C013", "2026-08-29", 75000, "repeat"])
    orders.append(["O102", "C021", "2026-08-11", 99000, None])
    note = workbook.create_sheet("MergedNote")
    note["A1"] = "Monthly summary"
    note.merge_cells("A1:B1")
    workbook.save(path)
    workbook.close()
    _canonicalize_zip(path)


def create_mini_dataset(root: Path) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    (root / "docs").mkdir(exist_ok=True)
    (root / "archive").mkdir(exist_ok=True)
    (root / "scans").mkdir(exist_ok=True)
    (root / "tables").mkdir(exist_ok=True)

    contact = (
        "담당자 주민번호 900101-1234567\n"
        "전화 010-1234-5678\n"
        "이메일 owner@example.com\n"
        "계좌 후보 110-123-456789\n"
    )
    (root / "docs" / "contact.txt").write_text(contact, encoding="utf-8")
    shutil.copyfile(root / "docs" / "contact.txt", root / "archive" / "contact_copy.txt")
    (root / "docs" / "return_policy_v1.txt").write_text("구버전 반품 가능 기간은 14일입니다.\n", encoding="utf-8")
    (root / "docs" / "return_policy_final_2026.txt").write_text("현재 반품 가능 기간은 30일입니다.\n", encoding="utf-8")
    _write_docx(root / "docs" / "warehouse_guide.docx")
    _write_text_pdf(root / "docs" / "shipping_policy.pdf")
    _write_blank_pdf(root / "scans" / "invoice_scan.pdf")
    (root / "tables" / "customers.csv").write_text(
        "customer_id,customer_name,email,active\n"
        "C013,한빛상사,sales@hanbit.example,true\n"
        "C021,새봄마트,,false\n",
        encoding="utf-8",
    )
    _write_xlsx(root / "tables" / "orders.xlsx")
    (root / "unsupported.bin").write_bytes(b"AX-SCANNER-UNSUPPORTED\n")

    for path in root.rglob("*"):
        if path.is_file():
            os.utime(path, (FIXED_TIMESTAMP, FIXED_TIMESTAMP))
    return root


def main() -> int:
    parser = argparse.ArgumentParser(description="Create the tiny deterministic 9/19 scanner dataset")
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    create_mini_dataset(args.output)
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
