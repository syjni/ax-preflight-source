#!/usr/bin/env python3
"""Validate Ceiling benchmark integrity without executing Kiro.

The deterministic plan replay below is a low-level tool-contract check.  It is
not evidence of Kiro or model runtime feasibility.
"""

from __future__ import annotations

import argparse
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any

from ax_agent.tools import AgentToolLayer
from ax_mcp.adapter import AxMcpAdapter
from ax_mcp.telemetry import InvocationLogger

from .generate_ceiling_dataset import AS_OF_DATE, DATASET_ID, SOURCE_RELATIVE
from .independent_ceiling_truth import independently_compute_truth
from .runtime_prompt import construct_runtime_prompt, project_runtime_prompt


ROOT = Path(__file__).resolve().parents[1]
REQUIRED_TASK_FIELDS = {
    "task_id", "category", "question", "expected_answer", "scoring_method",
    "required_sources", "primary_blocker", "control", "tool_plan",
}
VALID_CATEGORIES = {"knowledge", "operations", "cross_file"}
VALID_SCORING_METHODS = {"numeric", "exact", "set", "llm_judge"}
FORBIDDEN_RUNTIME_MARKERS = (
    "expected_answer", "primary_blocker", "tool_plan", "defect ground truth",
    "defect_ground_truth", "required_sources",
)


def _sha256(path: Path) -> str:
    import hashlib
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _check(name: str, passed: bool, detail: Any) -> dict[str, Any]:
    return {"name": name, "passed": passed, "detail": detail}


def _matches(actual: Any, expected: Any) -> bool:
    if isinstance(expected, float):
        try:
            return abs(float(actual) - expected) <= max(1e-9, abs(expected) * 1e-9)
        except (TypeError, ValueError):
            return False
    if isinstance(expected, int) and not isinstance(expected, bool):
        try:
            return int(float(actual)) == expected and float(actual).is_integer()
        except (TypeError, ValueError):
            return False
    return actual == expected


def _resolve_arguments(arguments: dict[str, Any], values: dict[str, Any]) -> dict[str, Any]:
    resolved = json.loads(json.dumps(arguments))
    for clause in resolved.get("filters", []):
        if "value_from" in clause:
            clause["value"] = values[clause.pop("value_from")]
    return resolved


def _table_from_search(results: list[Any], expected_file_id: str, columns: list[str]) -> str:
    for hit in results:
        if hit.document_id != expected_file_id:
            continue
        for table in hit.tables:
            if set(columns) <= set(table.column_names):
                return table.table_id
    raise ValueError("source table was not discovered with required columns")


def _table_from_scanner(layer: AgentToolLayer, expected_file_id: str, columns: list[str]) -> str:
    """Deterministic harness fallback; this is not a runtime discovery path."""
    for table in layer.report.tables:
        if table.file_id == expected_file_id and set(columns) <= set(table.column_names):
            return table.table_id
    raise ValueError("scanner report lacked a table with required columns")


def _validate_hidden_plan(layer: AgentToolLayer, task: dict[str, Any], file_ids: dict[str, str], expected: Any) -> dict[str, Any]:
    """Replay an evaluation-only plan against the four local tool contracts."""
    session = layer.start_task(task["task_id"], task["category"])
    discovered_documents: set[str] = set()
    discovered_tables: dict[str, str] = {}
    values: dict[str, Any] = {}
    answers: list[Any] = []
    for step in task["tool_plan"]["steps"]:
        tool = step["tool"]
        if tool == "search_documents":
            result = session.search_documents(step["query"], top_k=10)
            source = step["required_source"]
            target = file_ids[source]
            hit = next((item for item in result.results if item.document_id == target), None)
            if hit is None:
                # This harness explicitly permits scanner metadata fallback below.
                # It records contract behavior only and must not be called runtime evidence.
                pass
            discovered_documents.add(source)
            if hit is not None:
                for table in hit.tables:
                    discovered_tables.setdefault(source, table.table_id)
        elif tool == "read_document":
            source = step["source"]
            if source not in discovered_documents:
                raise ValueError("read document was not preceded by source discovery")
            read = session.read_document(file_ids[source])
            if not read.content.strip():
                raise ValueError("discovered document was empty")
        elif tool == "lookup_value":
            source = step["source"]
            if source not in discovered_documents:
                raise ValueError("lookup was not preceded by source discovery")
            try:
                table_id = _table_from_search(
                    session.search_documents(" ".join(step["required_columns"]), top_k=10).results,
                    file_ids[source], step["required_columns"],
                )
            except ValueError:
                table_id = _table_from_scanner(layer, file_ids[source], step["required_columns"])
            lookup = session.lookup_value(table_id, step["match_column"], step["value"], step["return_column"])
            if lookup.value is None or lookup.ambiguous:
                raise ValueError("lookup did not resolve one source value")
            values[step["save_as"]] = lookup.value
            discovered_tables[source] = table_id
        elif tool == "query_table":
            source = step["source"]
            if source not in discovered_documents:
                raise ValueError("table query was not preceded by source discovery")
            try:
                table_id = _table_from_search(
                    session.search_documents(" ".join(step["required_columns"]), top_k=10).results,
                    file_ids[source], step["required_columns"],
                )
            except ValueError:
                table_id = _table_from_scanner(layer, file_ids[source], step["required_columns"])
            result = session.query_table(table_id, **_resolve_arguments(step["arguments"], values))
            answers.extend(value for row in result.rows for value in row.values())
            discovered_tables[source] = table_id
        else:
            raise ValueError(f"unsupported plan tool: {tool}")
    if not task.get("expects_abstention") and task["category"] != "knowledge":
        if not any(_matches(value, expected) for value in answers):
            raise ValueError("tool result did not contain independently computed answer")
    return {"task_id": task["task_id"], "tool_calls": session.metrics().tool_calls, "passed": True}


def _zip_aware_text(path: Path) -> list[str]:
    if path.suffix.lower() not in {".docx", ".xlsx"}:
        return [path.read_bytes().decode("utf-8", errors="ignore")]
    chunks = []
    with zipfile.ZipFile(path) as archive:
        for member in archive.infolist():
            if member.is_dir():
                continue
            chunks.append(archive.read(member.filename).decode("utf-8", errors="ignore"))
    return chunks


def _strings(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        return [item for nested in value.values() for item in _strings(nested)]
    if isinstance(value, list):
        return [item for nested in value for item in _strings(nested)]
    return []


def _runtime_leakage(root: Path, source_root: Path, scan_report: Path, tasks: list[dict[str, Any]]) -> dict[str, Any]:
    agent_config = root / ".kiro" / "agents" / "ax-evaluation.json"
    runtime_paths = [path for path in source_root.rglob("*") if path.is_file()] + [scan_report]
    if agent_config.is_file():
        runtime_paths.append(agent_config)
    violations: list[dict[str, str]] = []
    for path in runtime_paths:
        for chunk in _zip_aware_text(path):
            folded = chunk.casefold()
            for marker in FORBIDDEN_RUNTIME_MARKERS:
                if marker.casefold() in folded:
                    violations.append({"path": str(path.relative_to(root)), "marker": marker})

    config = _read_json(agent_config) if agent_config.is_file() else None
    if config is not None and (config.get("resources") != [] or config.get("includeMcpJson") is not False):
        violations.append({"path": str(agent_config.relative_to(root)), "marker": "kiro_resource_boundary"})

    report = _read_json(scan_report)
    if any(table.get("sheet_name") == "SearchGuide" for table in report.get("tables", [])):
        violations.append({"path": scan_report.name, "marker": "benchmark_search_hint"})

    adapter = AxMcpAdapter(scan_report, source_root, InvocationLogger(stream=io.StringIO()))
    mcp_output = json.dumps({
        "schemas": adapter.list_tools(),
        "search": adapter.call_tool("search_documents", {"query": "거래처 주문 재고", "top_k": 5}),
    }, ensure_ascii=False).casefold()
    for marker in FORBIDDEN_RUNTIME_MARKERS:
        if marker.casefold() in mcp_output:
            violations.append({"path": "MCP output", "marker": marker})

    blind = _read_json(root / "runtime_blind_samples.json")
    samples = blind.get("samples", [])
    if len(samples) != len(tasks):
        violations.append({"path": "runtime_blind_samples.json", "marker": "sample_count"})
    for task, sample in zip(tasks, samples):
        projected = project_runtime_prompt(task)
        if set(sample) != {"category", "question", "runtime_prompt"}:
            violations.append({"path": "runtime_blind_samples.json", "marker": "sample_field_boundary"})
            continue
        if {"category": sample["category"], "question": sample["question"]} != projected:
            violations.append({"path": "runtime_blind_samples.json", "marker": "projection_mismatch"})
        prompt = construct_runtime_prompt(projected)
        if sample["runtime_prompt"] != prompt:
            violations.append({"path": "runtime_blind_samples.json", "marker": "constructor_mismatch"})
        forbidden_values = _strings(task.get("expected_answer")) + [task["primary_blocker"]]
        forbidden_values += task["required_sources"]
        for step in task["tool_plan"]["steps"]:
            forbidden_values += _strings(step.get("required_columns", []))
            forbidden_values += _strings(step.get("arguments", {}).get("filters", []))
            forbidden_values += _strings({key: step[key] for key in ("match_column", "return_column", "save_as") if key in step})
        for value in {item for item in forbidden_values if item and item not in projected.values() and item.casefold() not in projected["question"].casefold()}:
            if value.casefold() in prompt.casefold():
                violations.append({"path": "runtime prompt", "marker": "benchmark_value"})
                break
    return {
        "passed": not violations,
        "zip_aware": True,
        "runtime_file_count": len(runtime_paths),
        "violations": violations,
        "kiro_resource_boundary": {"agent_config_available": config is not None, "resources_empty": None if config is None else config.get("resources") == [], "include_mcp_json_disabled": None if config is None else config.get("includeMcpJson") is False},
    }


def _family_id(relative_path: str) -> str:
    path = Path(relative_path)
    normalized = re.sub(r"[^0-9a-z가-힣]+", "", path.stem.casefold())
    normalized = re.sub(r"(?:copy|duplicate|final|v\d+|수정)$", "", normalized)
    return f"{path.parent.as_posix()}/{normalized}"


def validate_ceiling_benchmark(root: Path = ROOT) -> dict[str, Any]:
    root = root.resolve()
    source_root = root / SOURCE_RELATIVE
    paths = {
        "manifest": root / "dataset_manifest.json", "benchmark": root / "benchmark_tasks.json",
        "truth": root / "ground_truth.json", "controls": root / "control_sources.json",
        "defects": root / "task_to_defect_manifest.json", "scan": root / "ceiling_scan_report.json",
        "provenance": root / "dataset_provenance.md", "blind_schema": root / "runtime_blind_sample.schema.json",
    }
    manifest = _read_json(paths["manifest"])
    benchmark = _read_json(paths["benchmark"])
    truth = _read_json(paths["truth"])
    controls = _read_json(paths["controls"])
    defects = _read_json(paths["defects"])
    tasks = benchmark["tasks"]
    checks: list[dict[str, Any]] = []

    checks.append(_check("ceiling_dataset_shape", manifest.get("dataset_id") == DATASET_ID and manifest.get("condition") == "ceiling" and manifest.get("file_count", 0) >= 20, {"file_count": manifest.get("file_count"), "format_counts": manifest.get("format_counts")}))
    required_formats = {"txt", "pdf", "docx", "csv", "xlsx"}
    checks.append(_check("supported_p0_formats_present", required_formats <= set(manifest.get("format_counts", {})), {"observed": manifest.get("format_counts")}))
    manifest_mismatches = [entry["relative_path"] for entry in manifest["files"] if not (source_root / entry["relative_path"]).is_file() or _sha256(source_root / entry["relative_path"]) != entry["sha256"]]
    checks.append(_check("dataset_manifest_integrity", not manifest_mismatches, {"mismatches": manifest_mismatches}))

    schema_errors = []
    categories: dict[str, int] = {}
    scoring_counts: dict[str, int] = {}
    seen = set()
    for task in tasks:
        categories[task.get("category", "")] = categories.get(task.get("category", ""), 0) + 1
        scoring = task.get("scoring_method", {})
        kind = scoring.get("type") if isinstance(scoring, dict) else None
        scoring_counts[str(kind)] = scoring_counts.get(str(kind), 0) + 1
        if REQUIRED_TASK_FIELDS - set(task) or task.get("task_id") in seen or task.get("category") not in VALID_CATEGORIES or kind not in VALID_SCORING_METHODS:
            schema_errors.append(task.get("task_id"))
        elif kind == "numeric" and not isinstance(task.get("expected_answer"), (int, float)):
            schema_errors.append(task["task_id"])
        elif kind == "set" and (not isinstance(task.get("expected_answer"), list) or set(scoring.get("required_items", [])) != set(task["expected_answer"]) or not scoring.get("reject_negated_items")):
            schema_errors.append(task["task_id"])
        seen.add(task.get("task_id"))
    checks.append(_check("benchmark_schema_and_scoring", len(tasks) >= 15 and not schema_errors, {"task_count": len(tasks), "category_counts": categories, "scoring_counts": scoring_counts, "errors": schema_errors}))

    source_paths = {entry["relative_path"] for entry in manifest["files"]}
    missing = sorted({source for task in tasks for source in task["required_sources"]} - source_paths)
    checks.append(_check("required_source_validation", not missing, {"missing": missing}))

    order_file = source_root / "04_주문" / "주문_원장_2026.xlsx"
    from openpyxl import load_workbook
    workbook = load_workbook(order_file, read_only=True, data_only=True)
    try:
        order_rows = list(workbook["Orders"].iter_rows(values_only=True))
    finally:
        workbook.close()
    order_date_index = list(order_rows[0]).index("order_date")
    max_order_date = max(str(row[order_date_index]) for row in order_rows[1:])
    checks.append(_check("as_of_date_order_boundary", max_order_date <= AS_OF_DATE.isoformat(), {"as_of_date": AS_OF_DATE.isoformat(), "max_order_date": max_order_date}))

    independent = independently_compute_truth(source_root)
    independent_answers = independent["answers"]
    truth_by_id = {item["task_id"]: item["expected_answer"] for item in truth["tasks"]}
    truth_errors = [task["task_id"] for task in tasks if independent_answers.get(task["task_id"]) != task["expected_answer"] or truth_by_id.get(task["task_id"]) != independent_answers.get(task["task_id"])]
    checks.append(_check("independent_source_truth_verification", not truth_errors and truth.get("independent_method") == independent["method"], {"mismatched_tasks": truth_errors, "method": independent["method"]}))

    control_errors = []
    expected_controls = sorted({source for task in tasks if task["control"] for source in task["required_sources"]})
    if controls.get("protected_sources") != expected_controls:
        control_errors.append("protected_source_union")
    observed_hashes = {source: _sha256(source_root / source) for source in controls.get("protected_sources", []) if (source_root / source).is_file()}
    if observed_hashes != controls.get("protected_source_sha256"):
        control_errors.append("source_hashes")
    all_families: dict[str, list[str]] = {}
    for path in source_root.rglob("*"):
        if path.is_file():
            relative = path.relative_to(source_root).as_posix()
            all_families.setdefault(_family_id(relative), []).append(relative)
    sibling_conflicts = []
    for family in controls.get("protected_logical_families", []):
        observed = sorted(all_families.get(family["family_id"], []))
        if observed != family.get("allowed_members"):
            sibling_conflicts.append({"family_id": family["family_id"], "observed": observed})
    if sibling_conflicts:
        control_errors.append("duplicate_or_conflicting_siblings")
    checks.append(_check("control_source_hash_and_family_protection", not control_errors, {"protected_source_count": len(expected_controls), "errors": control_errors, "sibling_conflicts": sibling_conflicts, "stability_metric": controls.get("later_retrieval_result_stability_metric")}))

    control_ids = {task["task_id"] for task in tasks if task["control"]}
    defect_entries = defects.get("task_defects", [])
    defect_errors = []
    expected_sensitive = {task["task_id"] for task in tasks if not task["control"]}
    if {entry.get("task_id") for entry in defect_entries} != expected_sensitive:
        defect_errors.append("task_coverage")
    for entry in defect_entries:
        if not entry.get("primary_defect_id") or set(entry.get("affected_sources", [])) & set(controls.get("protected_sources", [])) or set(entry.get("protected_dependencies", [])) != set(controls.get("protected_sources", [])) or set(entry.get("control_task_ids_excluded", [])) != control_ids:
            defect_errors.append(entry.get("task_id"))
    checks.append(_check("task_to_defect_manifest", defects.get("status") == "PLANNING_ONLY_NO_DEFECTS_INJECTED" and not defect_errors, {"defect_sensitive_task_count": len(defect_entries), "errors": defect_errors}))

    layer = AgentToolLayer(paths["scan"], source_root)
    file_ids = {record.relative_path: record.file_id for record in layer.report.files}
    deterministic, deterministic_errors = [], []
    for task in tasks:
        try:
            deterministic.append(_validate_hidden_plan(layer, task, file_ids, independent_answers[task["task_id"]]))
        except Exception as exc:
            deterministic_errors.append({"task_id": task["task_id"], "error": f"{type(exc).__name__}: {exc}"})
    checks.append(_check("deterministic_tool_contract_validation", not deterministic_errors, {"scope": "low-level hidden-plan replay only; not Kiro/runtime feasibility", "validated_tasks": len(deterministic), "failures": deterministic_errors, "tool_call_counts": {item["task_id"]: item["tool_calls"] for item in deterministic}}))

    leakage = _runtime_leakage(root, source_root, paths["scan"], tasks)
    checks.append(_check("runtime_leakage_prevention", leakage["passed"], leakage))
    checks.append(_check("blind_runtime_protocol_schema_present", paths["blind_schema"].is_file(), {"execution_status": "NOT_RUN"}))
    checks.append(_check("provenance_present", paths["provenance"].is_file() and "No public dataset was incorporated" in paths["provenance"].read_text(encoding="utf-8"), {"path": paths["provenance"].name}))

    passed = all(check["passed"] for check in checks)
    return {
        "schema_version": "ax-ceiling-benchmark-validation-v2",
        "dataset_id": DATASET_ID,
        "passed": passed,
        "safety_status": "NOT_SAFE",
        "safety_scope": "Static Ceiling-data integrity passed, but this artifact is not authorization to begin the 9/22 runtime experiment.",
        "remaining_weaknesses": [
            "Kiro was intentionally not run; the blind runtime sample protocol remains NOT_RUN.",
            "Deterministic tool validation is a low-level scanner-metadata harness, not Kiro/runtime retrieval feasibility evidence.",
        ],
        "checks": checks,
        "summary": {"task_count": len(tasks), "category_counts": categories, "control_task_count": len(control_ids), "scoring_counts": scoring_counts, "deterministic_tool_validation_count": len(deterministic)},
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate AX Ceiling benchmark integrity without executing Kiro")
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = validate_ceiling_benchmark(args.root)
    output = args.output or (args.root / "ceiling_validation.json")
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(output), "passed": result["passed"]}, ensure_ascii=False))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
