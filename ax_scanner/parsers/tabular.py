from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from ax_scanner.profiling import profile_rows

from .base import ParsedContent


def _read_csv(path: Path) -> tuple[list[list[str]], str, str]:
    last_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp949"):
        try:
            text = path.read_text(encoding=encoding)
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
                delimiter = dialect.delimiter
            except csv.Error:
                delimiter = ","
            return list(csv.reader(text.splitlines(), delimiter=delimiter)), encoding, delimiter
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Unable to decode CSV with utf-8 or cp949: {last_error}")


def parse_csv(path: Path, file_id: str, **_: object) -> ParsedContent:
    rows, encoding, delimiter = _read_csv(path)
    table = profile_rows(rows, file_id=file_id, sheet_name="CSV")
    text = "\n".join(" | ".join(str(value) for value in row) for row in rows)
    return ParsedContent(
        parser="csv",
        text=text,
        tables=[table],
        metadata={"encoding": encoding, "delimiter": delimiter},
    )


def parse_xlsx(path: Path, file_id: str, **_: object) -> ParsedContent:
    workbook = load_workbook(filename=str(path), read_only=False, data_only=True)
    tables = []
    text_parts = []
    try:
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            merged = sorted(str(cell_range) for cell_range in sheet.merged_cells.ranges)
            tables.append(profile_rows(rows, file_id=file_id, sheet_name=sheet.title, merged_ranges=merged))
            for row in rows:
                values = [str(value) for value in row if value is not None and str(value).strip()]
                if values:
                    text_parts.append(" | ".join(values))
    finally:
        workbook.close()
    return ParsedContent(
        parser="xlsx",
        text="\n".join(text_parts),
        tables=tables,
        metadata={"sheet_count": len(tables)},
    )

