from __future__ import annotations

from pathlib import Path

from docx import Document

from .base import ParsedContent


def parse_docx(path: Path, file_id: str, **_: object) -> ParsedContent:
    document = Document(str(path))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                parts.append(" | ".join(values))
    return ParsedContent(
        parser="docx",
        text="\n".join(parts),
        metadata={"paragraph_count": len(document.paragraphs), "table_count": len(document.tables)},
    )

